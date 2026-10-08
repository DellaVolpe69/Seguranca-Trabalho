"""Utilitários de tela compartilhados pelas páginas."""

import re
from datetime import date, datetime, time

import pandas as pd
import streamlit as st

OUTRO = "OUTRO — digitar"
VAZIO = "—"

# Funções da RQ 05 (aba "Listas"). Viram as opções do campo FUNÇÃO, junto
# com as que já foram usadas no banco — quem lança escolhe em vez de digitar.
FUNCOES_RQ05 = [
    "Ajudante Manutenção Predial", "Líder de Manutenção Predial", "Auxiliar Manutenção Predial",
    "Auxiliar Almoxarifado", "Supervisor de Almoxarifado", "Auxiliar de Borracheiro",
    "Borracheiro", "Carpinteiro", "Eletricista auto", "Apontador", "Sup. de Manutenção",
    "Gerente de Manutenção", "Lavador de Veículos", "Chefe de Manutenção",
    "Auxiliar de Manutenção Geral", "Mecânico de Carreta PL", "Mecânico Diesel PL / SR",
    "Mecânico Líder", "Operador de Check list", "Pintor Líder", "Polidor", "Frentista",
    "Soldador", "Motorista", "Ajudante de Cargas", "Conferente", "Supervisor de Armazém",
    "Gerente de Armazém", "Líder de Motorista", "Operador de Empilhadeira", "Auxiliar Operacional",
]
# Sim/Não com terceira opção: "não informado" fica NULL, não vira "Não".
SIM_NAO = {VAZIO: None, "Sim": True, "Não": False}


# ---------------------------------------------------------------------
# Conversões
# ---------------------------------------------------------------------

def vazio(valor) -> bool:
    if valor is None:
        return True
    if isinstance(valor, str):
        return not valor.strip()
    try:
        return bool(pd.isna(valor))
    except (TypeError, ValueError):
        return False


def texto(valor):
    """Texto limpo, ou None se vazio."""
    if vazio(valor):
        return None
    return re.sub(r"\s+", " ", str(valor)).strip()


def para_data(valor):
    if vazio(valor):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    return pd.to_datetime(valor).date()


def para_hora(valor):
    """'14:30:00' (como o Supabase devolve) -> time(14, 30)."""
    if vazio(valor):
        return None
    if isinstance(valor, time):
        return valor
    return time.fromisoformat(str(valor)[:8])


def para_numero(valor):
    return None if vazio(valor) else float(valor)


def rotulo_sim_nao(valor) -> str:
    if vazio(valor):
        return VAZIO
    return "Sim" if bool(valor) else "Não"


def fmt_data(valor) -> str:
    d = para_data(valor)
    return d.strftime("%d/%m/%Y") if d else ""


# ---------------------------------------------------------------------
# CPF
# ---------------------------------------------------------------------

def so_digitos(valor) -> str:
    return re.sub(r"\D", "", "" if vazio(valor) else str(valor))


def cpf_valido(cpf: str) -> bool:
    if not re.fullmatch(r"\d{11}", cpf or "") or cpf == cpf[0] * 11:
        return False
    for tamanho in (9, 10):
        soma = sum(int(cpf[i]) * (tamanho + 1 - i) for i in range(tamanho))
        digito = soma * 10 % 11 % 10
        if digito != int(cpf[tamanho]):
            return False
    return True


def erro_cpf(cpf_digitado) -> str | None:
    """Mensagem para quem digitou, ou None se o CPF é válido."""
    cpf = so_digitos(cpf_digitado)
    if not cpf:
        return "Informe o CPF."
    if len(cpf) != 11:
        return f"O CPF precisa ter 11 números — o digitado tem {len(cpf)}."
    if not cpf_valido(cpf):
        return (f"O CPF {fmt_cpf(cpf)} não existe: os 2 últimos números (dígitos verificadores) "
                "não batem com os 9 primeiros. Provavelmente algum número foi digitado errado.")
    return None


def fmt_cpf(cpf) -> str:
    d = so_digitos(cpf)
    return f"{d[:3]}.{d[3:6]}.{d[6:9]}-{d[9:]}" if len(d) == 11 else d


# ---------------------------------------------------------------------
# Mensagens que sobrevivem ao st.rerun()
# ---------------------------------------------------------------------

def guardar_msg(chave: str, tipo: str, texto_msg: str) -> None:
    st.session_state[chave] = (tipo, texto_msg)


def render_msg(chave: str) -> None:
    tipo, texto_msg = st.session_state.pop(chave, (None, None))
    if tipo == "success":
        st.success(texto_msg)
    elif tipo == "warning":
        st.warning(texto_msg)
    elif tipo == "error":
        st.error(texto_msg)


def mostrar_erros(erros: list, titulo: str = "Corrija antes de salvar:") -> None:
    st.error(f"{titulo}\n\n" + "\n".join(f"- {e}" for e in erros))


# ---------------------------------------------------------------------
# Campos
# ---------------------------------------------------------------------

def garantir_colunas(df: pd.DataFrame, colunas: list) -> pd.DataFrame:
    """Acrescenta (vazia) qualquer coluna esperada que o banco não devolveu."""
    faltando = [c for c in colunas if c not in df.columns]
    return df.assign(**{c: None for c in faltando}) if faltando else df


def opcoes_existentes(df: pd.DataFrame, coluna: str) -> list:
    """Valores já usados na coluna — viram as opções do selectbox."""
    if df.empty or coluna not in df:
        return []
    valores = {texto(v) for v in df[coluna]}
    return sorted(v for v in valores if v)


def campo_com_outro(rotulo: str, opcoes: list, key: str, valor_atual=None, ajuda=None):
    """Selectbox com os valores já usados + "OUTRO — digitar".

    Evita que "Filial SP", "filial sp" e "SP" virem três filiais no Power BI.
    """
    valor_atual = texto(valor_atual)
    lista = [VAZIO] + list(opcoes) + [OUTRO]
    if valor_atual is None:
        indice = 0
    elif valor_atual in opcoes:
        indice = lista.index(valor_atual)
    else:
        indice = len(lista) - 1

    escolha = st.selectbox(rotulo, lista, index=indice, key=key, help=ajuda)
    if escolha == OUTRO:
        inicial = valor_atual if valor_atual not in opcoes else ""
        return texto(st.text_input(f"{rotulo} (novo)", value=inicial or "", key=f"{key}_txt"))
    return None if escolha == VAZIO else escolha


def campo_sim_nao(rotulo: str, key: str, valor_atual=None):
    lista = list(SIM_NAO)
    escolha = st.selectbox(rotulo, lista, index=lista.index(rotulo_sim_nao(valor_atual)), key=key)
    return SIM_NAO[escolha]


def campo_lista(rotulo: str, opcoes: list, key: str, valor_atual=None):
    """Selectbox de lista fixa com opção vazia."""
    lista = [VAZIO] + list(opcoes)
    valor_atual = texto(valor_atual)
    indice = lista.index(valor_atual) if valor_atual in opcoes else 0
    escolha = st.selectbox(rotulo, lista, index=indice, key=key)
    return None if escolha == VAZIO else escolha


# ---------------------------------------------------------------------
# Extração em Excel
# ---------------------------------------------------------------------

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def planilha_excel(df: pd.DataFrame, colunas: dict, aba: str = "Dados") -> bytes:
    """A tabela da tela em .xlsx: mesmos títulos de coluna (o column_config do
    st.dataframe), data como data, Sim/Não no lugar de True/False, cabeçalho
    fixo e filtro em cada coluna."""
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    def rotulo(coluna):
        c = colunas.get(coluna, coluna)
        return c if isinstance(c, str) else (c.get("label") or coluna)

    def celula(v):
        if v is None or (not isinstance(v, (list, dict)) and pd.isna(v)):
            return None
        if isinstance(v, pd.Timestamp):
            return v.to_pydatetime().date() if v == v.normalize() else v.to_pydatetime()
        if hasattr(v, "item"):  # numpy → tipo do Python
            v = v.item()
        if isinstance(v, bool):
            return "Sim" if v else "Não"
        if isinstance(v, time):
            return v.strftime("%H:%M")
        return v

    wb = Workbook()
    ws = wb.active
    ws.title = aba[:31]
    ws.append([rotulo(c) for c in df.columns])
    for linha in df.itertuples(index=False):
        ws.append([celula(v) for v in linha])

    for cel in ws[1]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="E4610A")
        cel.alignment = Alignment(vertical="center")
    for i, coluna in enumerate(df.columns, start=1):
        letra = get_column_letter(i)
        maior = max([len(str(ws.cell(1, i).value))] + [len(str(v)) for v in df[coluna].head(500) if not vazio(v)])
        ws.column_dimensions[letra].width = min(max(maior + 2, 10), 60)
        for (cel,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
            if isinstance(cel.value, datetime):
                cel.number_format = "DD/MM/YYYY HH:MM"
            elif isinstance(cel.value, date):
                cel.number_format = "DD/MM/YYYY"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    saida = BytesIO()
    wb.save(saida)
    return saida.getvalue()


def baixar_excel(df: pd.DataFrame, colunas: dict, arquivo: str, key: str, aba: str = "Dados") -> None:
    """Botão "Baixar Excel" com a tabela filtrada da tela. `arquivo` sem extensão."""
    st.download_button("⬇️ Baixar Excel", planilha_excel(df, colunas, aba), file_name=f"{arquivo}.xlsx",
                       mime=XLSX, key=key)
